#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Из сырой выгрузки ESPON HOUSE4ALL — компактный CSV по муниципалитетам.

    python3 research/housing-europe/src/extract.py <папка с h4a_all.json и lau2021.xlsx>

Откуда сырьё (в репозиторий не коммитится, 61 МБ):
  * https://gis-server.espon.eu/arcgis/rest/services/Hosted/HOUSE4ALL_data_at_LAU_level/FeatureServer/0
    — публичный Feature Service, из которого рисует карты StoryMap проекта ESPON
    HOUSE4ALL. Снят 14 сентября 2026 постранично (query, 2000 записей за раз,
    returnGeometry=false, returnCentroid=true, outSR=4326), 89 601 запись.
    Сервер закрыт Cloudflare-проверкой для curl, выгрузка шла из браузера.
  * https://ec.europa.eu/eurostat/documents/345175/501971/EU-27-LAU-2021-NUTS-2021.xlsx
    — таблица соответствия LAU 2021 Евростата: степень урбанизации (DEGURBA:
    1 город, 2 пригород и малый город, 3 село), признак прибрежной зоны и
    латинское написание имени (им заменяются кириллица и греческое письмо).
    Только ЕС-27: у Швейцарии, Норвегии, Исландии и Лихтенштейна этих признаков нет.

Что пишется: research/data/house4all-lau-2024.csv — одна строка на LAU. Цены по
пяти классам площади (sp_corr_a…e, rp_corr_a…e) в CSV не идут: страница считает
по общей скорректированной цене, как и сам ESPON (sa_m2 и ra_m2 сервиса
воспроизводятся из income24, sp_corr, rp_corr и rate точно, до 1e-9).
"""
import csv, json, os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..', '..'))
OUT = os.path.join(ROOT, 'research', 'data', 'house4all-lau-2024.csv')

COLS = ['mun_id', 'cc', 'name', 'nuts3', 'nuts2', 'pop', 'area_km2', 'income24',
        'sp_corr', 'sp_unc', 'rp_corr', 'rp_unc', 'rate', 'sa_m2', 'ra_m2',
        's_listings', 'r_listings', 'lat', 'lon', 'degurba', 'coastal', 'lau_year']


def degurba(src):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(src, 'lau2021.xlsx'), read_only=True)
    m = {}
    for name in wb.sheetnames:
        if len(name) != 2:
            continue
        for i, row in enumerate(wb[name].iter_rows(values_only=True)):
            if i == 0 or not row[1]:
                continue
            m[name + '_' + str(row[1])] = (row[7], 1 if row[8] == 'yes' else 0, row[3] or '')
    return m


# Греческие муниципалитеты в таблице Евростата под другими кодами, латинского
# имени для них не найти — транслитерация по ЕЛОТ 743, чтобы поиск находил Athina.
EL_MAP = {'α': 'a', 'ά': 'a', 'β': 'v', 'γ': 'g', 'δ': 'd', 'ε': 'e', 'έ': 'e', 'ζ': 'z', 'η': 'i', 'ή': 'i',
          'θ': 'th', 'ι': 'i', 'ί': 'i', 'ϊ': 'i', 'ΐ': 'i', 'κ': 'k', 'λ': 'l', 'μ': 'm', 'ν': 'n', 'ξ': 'x',
          'ο': 'o', 'ό': 'o', 'π': 'p', 'ρ': 'r', 'σ': 's', 'ς': 's', 'τ': 't', 'υ': 'y', 'ύ': 'y', 'ϋ': 'y',
          'ΰ': 'y', 'φ': 'f', 'χ': 'ch', 'ψ': 'ps', 'ω': 'o', 'ώ': 'o'}


def translit_el(s):
    out = []
    for ch in s:
        low = ch.lower()
        if low in EL_MAP:
            t = EL_MAP[low]
            out.append(t.capitalize() if ch.isupper() else t)
        else:
            out.append(ch)
    s = ''.join(out)
    for a, b in (('Pseydokoinotita ', ''), ('Dimotiki Koinotita ', ''), ('Koinotita ', ''), ('Dimos ', '')):
        s = s.replace(a, b)
    return s


def num(v, nd=None):
    if v is None:
        return ''
    if isinstance(v, float):
        return ('%.' + str(nd) + 'f') % v if nd is not None else repr(v)
    return v


def main():
    src = sys.argv[1]
    pages = json.load(open(os.path.join(src, 'h4a_all.json'), encoding='utf-8'))
    deg = degurba(src)
    rows = []
    for p in pages:
        for f in p['features']:
            a = f['attributes']; c = f.get('centroid') or {}
            d = deg.get(a['mun_id'], ('', '', ''))
            # Болгария, Греция и Кипр в сервисе названы кириллицей и греческим; для
            # поиска берётся латинское имя из таблицы Евростата, если оно есть.
            name = d[2] if d[2] and a['cntr_code'] in ('BG', 'EL', 'CY') else a['lau_name']
            if a['cntr_code'] == 'EL':
                name = translit_el(name)
            rows.append([
                a['mun_id'], a['cntr_code'], name, a['nuts3'], a['nuts2'],
                num(a['pop_2021'], 0), num(a['area_km2'], 2), num(a['income24'], 0),
                num(a['sp_corr'], 1), num(a['sp_unc'], 1), num(a['rp_corr'], 3), num(a['rp_unc'], 3),
                num(a['sale_mor_1'], 3), num(a['sa_m2'], 2), num(a['ra_m2'], 2),
                num(a['s_num_list'], 0), num(a['r_num_list'], 0),
                num(c.get('y'), 4), num(c.get('x'), 4), d[0], d[1], a['year']])
    rows.sort(key=lambda r: r[0])
    with open(OUT, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(COLS)
        w.writerows(rows)
    print('%s: %d строк' % (os.path.relpath(OUT, ROOT), len(rows)))


if __name__ == '__main__':
    main()
